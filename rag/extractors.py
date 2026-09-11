"""Turn source files into a flat list of located text blocks.

Every block remembers the page or slide it came from so answers can cite
"Lecture 2, slide 14" instead of pointing vaguely at a whole deck.
"""

from __future__ import annotations

import csv
import html.parser
import io
import re
from dataclasses import dataclass
from pathlib import Path

from rag.normalize import clean

SUPPORTED = {".pdf", ".pptx", ".docx", ".md", ".markdown", ".txt", ".html", ".htm",
             ".csv", ".vtt", ".srt"}


class ExtractionError(RuntimeError):
    pass


@dataclass
class Block:
    """A run of text plus where in the document it lives."""

    text: str
    page: int          # 1-based page / slide number; 0 when the format has no pages
    label: str = ""    # "page" or "slide"
    section: str = ""  # caller-supplied label; overrides the chunker's guess


def _pdf(path: Path) -> list[Block]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - depends on install
        raise ExtractionError("PDF support needs 'pypdf' (pip install pypdf)") from exc

    blocks: list[Block] = []
    reader = PdfReader(str(path))
    for i, page in enumerate(reader.pages, start=1):
        text = clean(page.extract_text() or "")
        if text:
            blocks.append(Block(text=text, page=i, label="page"))
    if not blocks:
        raise ExtractionError(
            f"No text found in {path.name}. If it is a scan, OCR it first."
        )
    return blocks


def _pptx(path: Path) -> list[Block]:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - depends on install
        raise ExtractionError("PPTX support needs 'python-pptx' (pip install python-pptx)") from exc

    blocks: list[Block] = []
    deck = Presentation(str(path))
    for i, slide in enumerate(deck.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            parts.extend(_shape_text(shape))
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            parts.append(f"[Speaker notes] {notes}")
        text = clean("\n".join(p for p in parts if p.strip()))
        if text:
            blocks.append(Block(text=text, page=i, label="slide"))
    if not blocks:
        raise ExtractionError(f"No text found in {path.name}.")
    return blocks


def _shape_text(shape) -> list[str]:
    """Recurse through groups and tables, which hold most of the content in decks."""
    out: list[str] = []
    if shape.shape_type == 6 and hasattr(shape, "shapes"):  # GROUP
        for child in shape.shapes:
            out.extend(_shape_text(child))
        return out
    if getattr(shape, "has_text_frame", False):
        out.append(shape.text_frame.text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
    return out


def _docx(path: Path) -> list[Block]:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - depends on install
        raise ExtractionError("DOCX support needs 'python-docx' (pip install python-docx)") from exc

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = clean("\n".join(parts))
    if not text:
        raise ExtractionError(f"No text found in {path.name}.")
    return [Block(text=text, page=0)]


class _HTMLText(html.parser.HTMLParser):
    _SKIP = {"script", "style", "head"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skipping = 0

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self._skipping += 1

    def handle_endtag(self, tag):
        if tag in self._SKIP and self._skipping:
            self._skipping -= 1
        elif tag in {"p", "div", "li", "h1", "h2", "h3", "h4", "br", "tr"}:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skipping:
            self.parts.append(data)


def _html(path: Path) -> list[Block]:
    parser = _HTMLText()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    text = clean("".join(parser.parts))
    if not text:
        raise ExtractionError(f"No text found in {path.name}.")
    return [Block(text=text, page=0)]


def _csv(path: Path) -> list[Block]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    rows = list(csv.reader(io.StringIO(raw)))
    text = clean("\n".join(" | ".join(r) for r in rows if any(c.strip() for c in r)))
    if not text:
        raise ExtractionError(f"No rows found in {path.name}.")
    return [Block(text=text, page=0)]


def _plain(path: Path) -> list[Block]:
    text = clean(path.read_text(encoding="utf-8", errors="replace"))
    if not text:
        raise ExtractionError(f"{path.name} is empty.")
    return [Block(text=text, page=0)]


# WebVTT / SubRip cue header, e.g. "00:12:30.500 --> 00:12:34.000".
_CUE = re.compile(
    r"(?:(\d{1,2}):)?(\d{1,2}):(\d{2})[.,](\d{1,3})\s*-->\s*"
    r"(?:(\d{1,2}):)?(\d{1,2}):(\d{2})[.,](\d{1,3})"
)
_CAPTION_NOISE = re.compile(r"^(WEBVTT|NOTE|STYLE|REGION)\b", re.I)
_TAG = re.compile(r"</?[cuivb][^>]*>|<\d{2}:\d{2}:\d{2}[.,]\d{1,3}>")

CAPTION_WINDOW_SECONDS = 90


def _timestamp(seconds: int) -> str:
    hours, rest = divmod(int(seconds), 3600)
    minutes, secs = divmod(rest, 60)
    if hours:
        return f"{hours}:{minutes:02d}:{secs:02d}"
    return f"{minutes}:{secs:02d}"


def _captions(path: Path) -> list[Block]:
    """Parse .vtt/.srt into blocks labelled with the time they were spoken.

    Cues are a line or two each, far too small to retrieve on their own, so
    they are pooled into windows. The label is the window's start time, which
    is what a student actually needs: somewhere to scrub to in the recording.
    """
    raw = path.read_text(encoding="utf-8-sig", errors="replace")

    cues: list[tuple[int, str]] = []
    start: int | None = None
    spoken: list[str] = []

    def flush() -> None:
        if start is not None and spoken:
            cues.append((start, " ".join(spoken)))

    for line in raw.splitlines():
        line = line.strip()
        match = _CUE.search(line)
        if match:
            flush()
            hours, minutes, secs = match.group(1) or 0, match.group(2), match.group(3)
            start = int(hours) * 3600 + int(minutes) * 60 + int(secs)
            spoken = []
            continue
        if not line or line.isdigit() or _CAPTION_NOISE.match(line):
            continue
        if start is not None:
            spoken.append(_TAG.sub("", line))
    flush()

    if not cues:
        raise ExtractionError(
            f"No caption cues found in {path.name}. Expected WebVTT or SubRip format."
        )

    blocks: list[Block] = []
    window_start = cues[0][0]
    buffer: list[str] = []
    previous = ""
    for at, text in cues:
        # Close the window before adding this cue, so a cue that starts a new
        # window is the first line in it and the timestamp label stays honest.
        if buffer and at - window_start >= CAPTION_WINDOW_SECONDS:
            body = clean(" ".join(buffer))
            if body:
                blocks.append(Block(text=body, page=0, section=_timestamp(window_start)))
            window_start, buffer = at, []
        # Rolling captions repeat the previous line as they scroll; drop those.
        if text and text != previous:
            buffer.append(text)
            previous = text
    body = clean(" ".join(buffer))
    if body:
        blocks.append(Block(text=body, page=0, section=_timestamp(window_start)))

    if not blocks:
        raise ExtractionError(f"{path.name} contained no caption text.")
    return blocks


_HANDLERS = {
    ".pdf": _pdf,
    ".pptx": _pptx,
    ".docx": _docx,
    ".md": _plain,
    ".markdown": _plain,
    ".txt": _plain,
    ".html": _html,
    ".htm": _html,
    ".csv": _csv,
    ".vtt": _captions,
    ".srt": _captions,
}


def extract(path: Path) -> list[Block]:
    """Read `path` into located, cleaned text blocks."""
    suffix = path.suffix.lower()
    handler = _HANDLERS.get(suffix)
    if handler is None:
        raise ExtractionError(
            f"Cannot read '{suffix or path.name}'. Supported: {', '.join(sorted(SUPPORTED))}"
        )
    return handler(path)
